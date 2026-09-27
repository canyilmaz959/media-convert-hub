import sys
import os
import gc
import tempfile

if len(sys.argv) < 4:
    print("ERROR: Eksik argümanlar. Kullanım: main.py <input> <output> <format>")
    sys.exit(1)

input_path = sys.argv[1]
output_path = sys.argv[2]
target_format = sys.argv[3].lower().strip()

# Render çözünürlüğü düşük tutuluyor çünkü Render servisi 512 MB RAM ile çalışıyor.
# Daha yüksek kalite gerekirse ortam değişkeni ile artırılabilir.
PDF_RENDER_DPI = int(os.getenv("PDF_RENDER_DPI", "96"))


def convert_to_docx(infile, outfile):
    from pdf2docx import Converter

    cv = Converter(infile)
    try:
        cv.convert(outfile)
    finally:
        cv.close()


def convert_pdf_pages_to_images(infile):
    """PDF sayfalarını tek tek işler; tüm sayfaları RAM'e yüklemez."""
    import fitz

    doc = fitz.open(infile)
    try:
        for page_number in range(len(doc)):
            page = doc[page_number]
            pix = page.get_pixmap(
                dpi=PDF_RENDER_DPI,
                colorspace=fitz.csRGB,
                alpha=False
            )
            yield page_number, pix
            del pix
            gc.collect()
    finally:
        doc.close()


def convert_to_pptx(infile, outfile):
    from pptx import Presentation

    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]

    # Her Python süreci için benzersiz geçici klasör.
    # Aynı anda birden fazla dönüşüm çalışırken dosyalar birbirine girmez.
    with tempfile.TemporaryDirectory(prefix="pdf_pptx_") as temp_dir:
        temp_path = os.path.join(temp_dir, "page.png")

        for page_number, pix in convert_pdf_pages_to_images(infile):
            pix.save(temp_path)

            slide = prs.slides.add_slide(blank_slide_layout)
            slide.shapes.add_picture(
                temp_path,
                0,
                0,
                width=prs.slide_width,
                height=prs.slide_height
            )

            print(f"Sayfa işlendi: {page_number + 1}")
            del pix
            gc.collect()

    prs.save(outfile)


def convert_to_odp(infile, outfile):
    from odf.opendocument import OpenDocumentPresentation
    from odf.draw import Page, Frame, Image

    doc = OpenDocumentPresentation()

    # odfpy, addPicture() ile verilen dosyayı doc.save() sırasında paketleyebilir.
    # Bu yüzden PNG'leri save() tamamlanana kadar silmiyoruz.
    # RAM'i sınırlamak için yine de sayfaları tek tek render ediyoruz.
    with tempfile.TemporaryDirectory(prefix="pdf_odp_") as temp_dir:
        for page_number, pix in convert_pdf_pages_to_images(infile):
            temp_path = os.path.join(temp_dir, f"page_{page_number + 1}.png")
            pix.save(temp_path)

            slide = Page(
                name=f"slide_{page_number + 1}",
                masterpagename="Standard"
            )
            doc.body.appendChild(slide)

            frame = Frame(
                width="28cm",
                height="21cm",
                x="0cm",
                y="0cm"
            )
            slide.appendChild(frame)

            href = doc.addPicture(temp_path)
            img_node = Image(href=href)
            frame.appendChild(img_node)

            print(f"Sayfa işlendi: {page_number + 1}")
            del pix
            gc.collect()

        # Geçici görseller bu noktaya kadar mevcut olmalı.
        doc.save(outfile)


def convert_to_potx(infile, outfile):
    from pptx import Presentation

    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory(prefix="pdf_potx_") as temp_dir:
        temp_path = os.path.join(temp_dir, "page.png")

        for page_number, pix in convert_pdf_pages_to_images(infile):
            pix.save(temp_path)

            slide = prs.slides.add_slide(blank_slide_layout)
            slide.shapes.add_picture(
                temp_path,
                0,
                0,
                width=prs.slide_width,
                height=prs.slide_height
            )

            print(f"Sayfa işlendi: {page_number + 1}")
            del pix
            gc.collect()

    prs.save(outfile)


try:
    if target_format == "docx":
        convert_to_docx(input_path, output_path)
    elif target_format == "pptx":
        convert_to_pptx(input_path, output_path)
    elif target_format == "odp":
        convert_to_odp(input_path, output_path)
    elif target_format == "potx":
        convert_to_potx(input_path, output_path)
    else:
        raise ValueError(f"Desteklenmeyen format: {target_format}")

    print("başarılı")
    sys.exit(0)

except Exception as e:
    print("başarısız hata:", str(e))
    sys.exit(1)
